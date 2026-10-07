# -*- coding: utf-8 -*-
r"""Formula-injection answer (Lead question 3).

1. Which writers reachable here turn a string into an evaluable cell?
   - openpyxl (the deliverable's writer)
   - XlsxWriter (bundled, second opinion: does it treat -, +, @ as formulas?)
2. Structural scan of a workbook produced by adif2xlsx.py for every formula
   artefact OOXML can carry: <f> elements, t="str" cached formula strings,
   defined names, external links.

Run: python verification\probe_injection.py
"""
from __future__ import annotations

import os
import re
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

PAYLOADS = ["=1+1", "-10", "+12", "@home", "- QRP note", " =1+1", "\t=1+1"]
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "out")


def openpyxl_probe():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    for i, v in enumerate(PAYLOADS, start=1):
        ws.cell(row=i, column=1, value=v)
    path = os.path.join(OUT, "_inj_openpyxl.xlsx")
    wb.save(path)
    xml = zipfile.ZipFile(path).read("xl/worksheets/sheet1.xml").decode()
    rows = {}
    for m in re.finditer(r"<c r=\"A(\d+)\"[^>]*>(.*?)</c>", xml):
        rows[int(m.group(1))] = m.group(2)
    print("openpyxl writer:")
    for i, v in enumerate(PAYLOADS, start=1):
        body = rows.get(i, "")
        kind = "FORMULA" if "<f>" in body or "<f " in body else "text"
        print(f"   {v!r:16} -> {kind:8} {body[:70]}")
    return path


def xlsxwriter_probe():
    try:
        import xlsxwriter
    except ImportError:
        print("XlsxWriter: not bundled")
        return None
    path = os.path.join(OUT, "_inj_xlsxwriter.xlsx")
    wb = xlsxwriter.Workbook(path)
    ws = wb.add_worksheet()
    for i, v in enumerate(PAYLOADS):
        ws.write(i, 0, v)
    wb.close()
    xml = zipfile.ZipFile(path).read("xl/worksheets/sheet1.xml").decode()
    rows = {}
    for m in re.finditer(r"<c r=\"A(\d+)\"[^>]*>(.*?)</c>", xml):
        rows[int(m.group(1))] = m.group(2)          # 1-based XML row = payload index
    print(f"\nXlsxWriter {xlsxwriter.__version__} writer (strings are shared-string indices):")
    for i, v in enumerate(PAYLOADS, start=1):
        body = rows.get(i, "")
        kind = "FORMULA" if "<f>" in body or "<f " in body else "text"
        print(f"   {v!r:16} -> {kind:8} {body[:70]}")
    return path


def package_scan(path: str) -> int:
    """Every formula artefact OOXML can carry, across the whole package."""
    problems = []
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        for name in names:
            if not name.endswith(".xml"):
                continue
            xml = z.read(name).decode("utf-8", errors="replace")
            if name.startswith("xl/worksheets/"):
                for m in re.finditer(r"<f[ >]", xml):
                    problems.append(f"{name}: <f> element at offset {m.start()}")
                for m in re.finditer(r'<c [^>]*t="str"', xml):
                    problems.append(f"{name}: cached formula string (t=\"str\")")
            if name == "xl/workbook.xml":
                for m in re.finditer(r"<definedName[^>]*>.*?</definedName>", xml, re.S):
                    text = m.group(0)
                    if "_xlnm._FilterDatabase" in text or "_FilterDatabase" in text:
                        continue                       # autofilter range: benign
                    problems.append(f"{name}: defined name {text[:80]}")
        for name in names:
            if "externalLink" in name:
                problems.append(f"{name}: external link part")
    print(f"\npackage scan of {os.path.basename(path)}: {len(problems)} formula artefact(s)")
    for p in problems[:10]:
        print("   ", p)
    return len(problems)


def main() -> None:
    os.makedirs(OUT, exist_ok=True)
    openpyxl_probe()
    xlsxwriter_probe()
    target = os.path.join(OUT, "q12.xlsx")          # produced by adif2xlsx.py
    if os.path.exists(target):
        package_scan(target)
    target2 = os.path.join(OUT, "q12b.xlsx")
    if os.path.exists(target2):
        package_scan(target2)


if __name__ == "__main__":
    main()

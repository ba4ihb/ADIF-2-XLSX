#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Tests for the web front end (src/webapp.py).

The browser is not driven directly; instead the service is started on a loopback
port and its HTTP API is exercised exactly as the page does.  That covers what
matters: the workbook the service produces must be the same one the command line
produces, the column selection must be honoured, and a local service must not be
a way to read files it was never given.

Run:  python tests/test_web.py
"""

from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(ROOT, "src")
SAMPLES = os.path.join(HERE, "fixtures", "samples")
REAL = os.path.join(HERE, "fixtures", "adi")
WEB = os.path.join(ROOT, "web")

sys.path.insert(0, SRC)
import adif2xlsx as core  # noqa: E402
import postage  # noqa: E402
import webapp  # noqa: E402

FAILURES = []
CHECKS = 0
OUT = tempfile.mkdtemp(prefix="adif2xlsx_web_suite_")


def check(condition, label, detail=""):
    global CHECKS
    CHECKS += 1
    if condition:
        print(f"  PASS  {label}")
    else:
        print(f"  FAIL  {label}" + (f"  [{detail}]" if detail else ""))
        FAILURES.append(label)


def read(path):
    with open(path, "rb") as fh:
        return fh.read()


def zip_of(folder_or_pairs):
    """A zip of files, as the page builds for a multi-file selection."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_STORED) as zf:
        for name, data in folder_or_pairs:
            zf.writestr(name, data)
    return buf.getvalue()


def scratch_dirs():
    """This service's own scratch folders, which must not survive a request."""
    return {n for n in os.listdir(tempfile.gettempdir())
            if n.startswith(("adif2xlsx_web_", "adif2xlsx_preview_"))}


class Service:
    def __init__(self):
        self.port = webapp.find_port()
        self.httpd = webapp.ThreadingHTTPServer(("127.0.0.1", self.port),
                                                webapp.Handler)
        threading.Thread(target=self.httpd.serve_forever, daemon=True).start()
        time.sleep(0.3)
        self.base = f"http://127.0.0.1:{self.port}"
        with urllib.request.urlopen(self.base + "/api/info") as response:
            self.info = json.load(response)
        self.token = self.info["token"]

    def get(self, path, raw=False):
        try:
            with urllib.request.urlopen(self.base + path) as response:
                body = response.read()
                return response.status, (body if raw else
                                         body.decode("utf-8", "replace")), \
                    dict(response.headers)
        except urllib.error.HTTPError as exc:
            return exc.code, exc.read(), dict(exc.headers)

    def post(self, path, body, headers=None):
        request = urllib.request.Request(self.base + path, data=body,
                                         method="POST")
        request.add_header("X-Adif2xlsx-Token", self.token)
        for key, value in (headers or {}).items():
            request.add_header(key, value)
        try:
            with urllib.request.urlopen(request) as response:
                return response.status, response.read(), dict(response.headers)
        except urllib.error.HTTPError as exc:
            return exc.code, exc.read(), dict(exc.headers)

    def stop(self):
        self.httpd.shutdown()
        self.httpd.server_close()


def report_of(headers):
    return json.loads(urllib.parse.unquote(
        headers.get("X-Adif2xlsx-Report", "%7B%7D")))


# ---------------------------------------------------------------------------
def test_static_and_info(svc):
    print("\n=== 页面与接口 ===")
    status, page, headers = svc.get("/")
    check(status == 200 and "ADIF" in page,
          "首页可以打开", f"{status} {len(page)}B")
    check("charset=utf-8" in headers.get("Content-Type", ""),
          "首页声明 UTF-8", headers.get("Content-Type", ""))
    check("开始转换" in page, "首页是中文界面")
    check("Drop" in page or "拖放" in page, "首页有拖放区域")

    status, css, _ = svc.get("/icon.png", raw=True)
    check(status == 200 and css[:4] == b"\x89PNG", "图标可以取到", str(status))

    check(svc.info["ok"] is True, "/api/info 返回 ok")
    # The required set matches the core's; the count is asserted through the
    # constant rather than a literal, because a literal here is a second source
    # of truth that drifts (it read 12 while three of those were dropped when
    # empty).
    check(len(svc.info["required"]) == len(core.REQUIRED_COLUMNS),
          "必选列数量与核心一致",
          f"{len(svc.info['required'])} vs {len(core.REQUIRED_COLUMNS)}")
    check([i["name"] for i in svc.info["required"]]
          == list(core.REQUIRED_COLUMNS),
          "必选列与核心同名同序")
    check(len(svc.info["optional"]) > 30, "附加列已列出",
          str(len(svc.info["optional"])))
    check(svc.info["mandatory"], "强制输出列已声明")
    # The output field starts EMPTY so the workbook is downloaded.  A local
    # path is something the user types, not something the tool presumes.
    check(svc.info["output"] == "",
          "默认不预设保存路径（留空即下载）", repr(svc.info["output"]))
    check(svc.info.get("download_name"), "给出了下载文件名",
          str(svc.info.get("download_name")))
    check(svc.info.get("suggested_output"), "提供一个可选的默认位置",
          str(svc.info.get("suggested_output")))
    check({"common", "all", "required"} <= set(svc.info["presets"]),
          "三个快捷预设都在", str(list(svc.info["presets"])))
    labelled = [f for f in svc.info["required"] + svc.info["optional"]
                if f["label"]]
    check(len(labelled) == len(svc.info["required"]) + len(svc.info["optional"]),
          "每个字段都有中文名")


def test_preview(svc):
    print("\n=== 读取字段（预览）===")
    sample = os.path.join(SAMPLES, "sample_log.adi")
    status, body, _ = svc.post("/api/preview", read(sample))
    data = json.loads(body)
    check(status == 200 and data["ok"], "预览成功", str(status))
    check(data["files"][0]["records"] == 17, "预览报告 17 条记录",
          str(data["files"][0]))
    check("GRIDSQUARE" in data["available"], "预览列出可用字段",
          str(data["available"][:8]))

    # A file that is not ADIF must be reported, not crash the service.
    status, body, _ = svc.post("/api/preview", b"not adif at all")
    data = json.loads(body)
    check(status == 200 and data["files"][0]["error"],
          "非 ADIF 文件被逐条报告", str(data["files"][0])[:80])


def test_convert_download(svc):
    print("\n=== 转换并下载 ===")
    sample = os.path.join(SAMPLES, "sample_log.adi")
    status, body, headers = svc.post("/api/convert", read(sample))
    check(status == 200, "转换返回 200", str(status))
    check("spreadsheetml" in headers.get("Content-Type", ""),
          "返回 xlsx 类型", headers.get("Content-Type", ""))
    check(body[:2] == b"PK", "返回值是 xlsx 容器")
    check("attachment" in headers.get("Content-Disposition", ""),
          "带下载文件名", headers.get("Content-Disposition", ""))

    report = report_of(headers)
    check(report.get("total") == 17, "报告 17 条记录", str(report.get("total")))
    check(report.get("columns", 0) >= 12, "报告输出列数",
          str(report.get("columns")))
    check("SOURCES" not in str(report), "报告不夹带无关内容")

    path = os.path.join(OUT, "download.xlsx")
    with open(path, "wb") as fh:
        fh.write(body)
    from openpyxl import load_workbook
    book = load_workbook(path)
    check("QSOs" in book.sheetnames and "Summary" in book.sheetnames,
          "工作簿含 QSOs 与 Summary", str(book.sheetnames))
    sheet = book["QSOs"]
    check(sheet.max_row - 2 == 17, "QSOs 有 17 行数据", str(sheet.max_row - 2))
    check([c.value for c in sheet[1]][:3] == ["CALL", "QSO_DATE", "TIME_ON_UTC"],
          "第一行是英文列名")
    check(sheet.cell(row=2, column=1).value == "对方呼号", "第二行是中文列名",
          str(sheet.cell(row=2, column=1).value))
    check(sheet.freeze_panes == "A3", "冻结两行表头", str(sheet.freeze_panes))

    # The Summary must explain the empty columns, as on the desktop.
    summary = book["Summary"]
    text = "\n".join(str(c.value) for row in summary.iter_rows()
                     for c in row if c.value)
    check("各列填充情况" in text, "Summary 含各列填充情况")
    check("国内" in text or "邮资" in text, "Summary 含邮资说明")


def test_multi_file_zip(svc):
    print("\n=== 多文件（zip 上传）===")
    pairs = [(name, read(os.path.join(SAMPLES, name)))
             for name in ("sample_log.adi", "sample_lotw_quirks.adi")]
    status, body, headers = svc.post("/api/convert", zip_of(pairs))
    check(status == 200, "多文件转换返回 200", str(status))
    report = report_of(headers)
    check(report.get("total") == 22, "两个文件合计 22 条",
          str(report.get("total")))
    check(len(report.get("sources", {})) == 2, "报告两个来源",
          str(report.get("sources")))
    check(len(report.get("column_names", [])) > 12,
          "多文件时附加列被合并", str(len(report.get("column_names", []))))

    # Only ADIF members are read out of the zip.
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_STORED) as zf:
        zf.writestr("readme.txt", "not adif")
    status, body, _ = svc.post("/api/convert", buf.getvalue())
    check(status == 400, "zip 里没有 ADIF 时返回 400", str(status))


def test_real_exports(svc):
    print("\n=== 六个真实导出 ===")
    pairs = [(name, read(os.path.join(REAL, name)))
             for name in sorted(os.listdir(REAL))]
    status, body, headers = svc.post("/api/convert", zip_of(pairs))
    check(status == 200, "转换成功", str(status))
    report = report_of(headers)
    check(report.get("total") == 249, "合计 249 条", str(report.get("total")))

    path = os.path.join(OUT, "real.xlsx")
    with open(path, "wb") as fh:
        fh.write(body)
    from openpyxl import load_workbook
    sheet = load_workbook(path)["QSOs"]
    header = [c.value for c in sheet[1]]
    entity_at = header.index("DXCC_ENTITY")
    postage_at = header.index("QSL_POSTAGE_AIR")
    rows = list(sheet.iter_rows(min_row=3, values_only=True))
    check(all(row[entity_at] for row in rows), "每行都有 DXCC 实体",
          str(sum(1 for row in rows if not row[entity_at])))
    check(all(row[postage_at] not in (None, "") for row in rows),
          "每行都有邮资",
          str(sum(1 for row in rows if row[postage_at] in (None, ""))))
    china = {row[postage_at] for row in rows if row[entity_at] == "China"}
    check(china == {1.2}, "中国大陆邮资是 1.2", str(china))
    # These exports happen to hold no Taiwanese station, so the rule is checked
    # through the API rather than through a fixture: mainland China and Taiwan
    # are different DXCC entities and must not share a postage figure.
    mainland = postage.quote("China").rates
    taiwan = postage.quote("Taiwan").rates
    check(mainland != taiwan, "大陆与台湾邮资不同",
          f"{mainland} vs {taiwan}")
    check(postage.quote("China").detail == "国内 1.20",
          "大陆邮资是单一数字 1.20", postage.quote("China").detail)
    regions = {row[entity_at] for row in rows}
    check("European Russia" in regions or "Asiatic Russia" in regions,
          "俄罗斯按实体拆分", str(sorted(regions))[:120])


def test_columns(svc):
    print("\n=== 指定列 ===")
    sample = os.path.join(SAMPLES, "sample_log.adi")
    chosen = ["CALL", "QSO_DATE", "MODE", "GRIDSQUARE", "MY_CITY"]
    status, body, headers = svc.post(
        "/api/convert", read(sample),
        {"X-Adif2xlsx-Columns": urllib.parse.quote(json.dumps(chosen))})
    check(status == 200, "指定列转换成功", str(status))
    path = os.path.join(OUT, "cols.xlsx")
    with open(path, "wb") as fh:
        fh.write(body)
    from openpyxl import load_workbook
    header = [c.value for c in load_workbook(path)["QSOs"][1]]
    check("GRIDSQUARE" in header, "指定的附加列被输出", str(header))
    check("MY_CITY" in header, "日志没有的列也按请求输出", str(header))
    check("QTH" not in header, "未指定的附加列不输出", str(header))
    check(all(c in header for c in ["CALL", "QSO_DATE", "MODE"]),
          "必选列仍在", str(header))

    # 仅必选列
    status, body, headers = svc.post(
        "/api/convert", read(sample),
        {"X-Adif2xlsx-Columns": urllib.parse.quote(json.dumps([]))})
    path = os.path.join(OUT, "required.xlsx")
    with open(path, "wb") as fh:
        fh.write(body)
    header = [c.value for c in load_workbook(path)["QSOs"][1]]
    check("GRIDSQUARE" not in header, "仅必选列时附加列不出现", str(header))
    check(all(c in header for c in webapp.core.MANDATORY_COLUMNS),
          "强制列始终存在",
          str([c for c in webapp.core.MANDATORY_COLUMNS if c not in header]))


def test_save_to_path(svc):
    print("\n=== 保存到指定路径 ===")
    sample = os.path.join(SAMPLES, "sample_log.adi")
    target = os.path.join(OUT, "saved.xlsx")
    status, body, _ = svc.post(
        "/api/convert", read(sample),
        {"X-Adif2xlsx-Output": urllib.parse.quote(target)})
    data = json.loads(body)
    check(status == 200 and data["ok"], "指定路径转换成功", str(status))
    check(os.path.exists(target), "文件确实写到该路径", target)
    check(data.get("saved_to") == target, "报告里带保存路径",
          str(data.get("saved_to")))

    # An extension-less path gets one, rather than writing an unnamed file.
    bare = os.path.join(OUT, "noext")
    status, body, _ = svc.post(
        "/api/convert", read(sample),
        {"X-Adif2xlsx-Output": urllib.parse.quote(bare)})
    check(os.path.exists(bare + ".xlsx"), "缺少扩展名时自动补 .xlsx",
          str(status))


def test_security(svc):
    print("\n=== 本地服务的安全边界 ===")
    sample = os.path.join(SAMPLES, "sample_log.adi")
    request = urllib.request.Request(svc.base + "/api/convert",
                                     data=read(sample), method="POST")
    request.add_header("X-Adif2xlsx-Token", "definitely-not-the-token")
    try:
        urllib.request.urlopen(request)
        check(False, "错误令牌必须被拒绝")
    except urllib.error.HTTPError as exc:
        check(exc.code == 403, "错误令牌返回 403", str(exc.code))

    request = urllib.request.Request(svc.base + "/api/convert",
                                     data=read(sample), method="POST")
    try:
        urllib.request.urlopen(request)
        check(False, "缺少令牌必须被拒绝")
    except urllib.error.HTTPError as exc:
        check(exc.code == 403, "缺少令牌返回 403", str(exc.code))

    # Traversal attempts must not reach anything outside web/.
    for probe in ("/../src/adif2xlsx.py", "/..%2Fsrc%2Fadif2xlsx.py",
                  "/../README.md", "/../../Windows/win.ini"):
        status, body, _ = svc.get(probe, raw=True)
        leaked = b"def parse_adif" in body or b"[fonts]" in body
        check(status in (400, 404) and not leaked,
              f"{probe} 读不到文件", f"{status} {body[:40]!r}")

    status, body, _ = svc.post("/api/convert", b"x" * 10)
    check(status in (400, 413), "过小/异常请求被拒绝", str(status))


def test_cleanup(svc):
    print("\n=== 不留下上传残留 ===")
    sample = os.path.join(SAMPLES, "sample_log.adi")
    before = scratch_dirs()
    svc.post("/api/convert", read(sample))
    svc.post("/api/preview", read(sample))
    time.sleep(0.2)
    leaked = scratch_dirs() - before
    check(not leaked, "服务不残留上传目录", str(sorted(leaked)))


def test_offline_assets():
    print("\n=== 页面不依赖外网 ===")
    page = open(os.path.join(WEB, "index.html"), encoding="utf-8").read()
    external = [token for token in ("http://", "https://", "//cdn", "cdnjs",
                                    "googleapis", "unpkg", "jsdelivr")
                if token in page]
    check(not external, "页面没有引用任何外部资源", str(external))
    check(os.path.isfile(os.path.join(WEB, "icon.png")), "图标文件存在")
    # The page must build a zip for multi-file uploads without a library.
    check("class ZipWriter" in page, "多文件打包是页面自带的")
    check("crc32" in page, "zip 校验和已实现")


def main() -> int:
    service = Service()
    try:
        test_static_and_info(service)
        test_preview(service)
        test_convert_download(service)
        test_multi_file_zip(service)
        test_real_exports(service)
        test_columns(service)
        test_save_to_path(service)
        test_security(service)
        test_cleanup(service)
        test_offline_assets()
    finally:
        service.stop()

    print(f"\n{'=' * 62}")
    print(f"checks: {CHECKS}   failures: {len(FAILURES)}")
    for item in FAILURES:
        print(f"  - {item}")
    return 1 if FAILURES else 0


if __name__ == "__main__":
    raise SystemExit(main())
